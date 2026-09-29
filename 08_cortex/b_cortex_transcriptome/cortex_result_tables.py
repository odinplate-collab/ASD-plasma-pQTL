# Assemble result tables (detection in brain cell types, complement cascade, disorders, C3 and IGFBP5 across layers) from the cortex, plasma and single-nucleus results.
# Inputs (relative to ASD_ROOT):
#   brain/data/asd_cortex/Gandal2022_MOESM5.xlsx            Gandal et al. 2022, Supplementary Data 3
#   brain/data/asd_cortex/gandal2022_complement_genes.tsv   (gandal2022_primary.py)
#   brain/data/asd_cortex/gandal2022_modules.tsv            (gandal2022_primary.py)
#   brain/data/replication/gandal2018a/TableS1_Microarray_MetaAnalyses.csv   Gandal et al. 2018, Table S1
#   brain/data/replication/replication_modules.tsv          (replication_modules.py)
#   brain/data/replication/zhang2023/pnas.2206758120.sd01.xlsx   Zhang et al. 2023, Dataset S1 (STG bulk, DESeq2)
#   brain/data/replication/zhang2023/pnas.2206758120.sd06.xlsx   Zhang et al. 2023, Dataset S6 (LCM neurons, DESeq2)
#   brain/data/brain_proteome/abraham2019_protein_table.tsv Abraham et al. 2019 Sci Rep, BA19 protein table
#     taken from the Supplementary Information (gene, p, fdr, val = ASD abundance as % of control)
#   panel12/dia_validation/stats_dep_concordance.csv           plasma TMT log2FC and P (06_DIA_validation/dia_validation.R)
#   brain/data/lr_singlecell/panel12_by_celltype_both_datasets.tsv   (../c_single_nucleus/sc_de_both.R)
#   brain/data/lr_singlecell/de_genes_both_datasets.tsv              (../c_single_nucleus/sc_de_both.R)
#   brain/data/lr_singlecell/panel12_meta_two_datasets.tsv           (../c_single_nucleus/panel12_meta.py)
# Outputs (relative to ASD_ROOT), all in brain/results/:
#   brain_cell_detection.tsv          detection and mean CPM of the 12 genes per dataset and cell type
#   brain_cell_detection_mean.tsv     % nuclei detecting each gene, mean of the two datasets
#   complement_cascade_genes.tsv        complement genes: RNA-seq and microarray log2FC / FDR
#   complement_cascade_layers.tsv       complement layers: median log2FC and competitive P
#   disorders_panel12.tsv     12 genes and complement module across disorders
#   c3_across_layers.tsv            C3 across plasma, cortex RNA, cortex protein, microglia RNA
#   gandal2022_regions.tsv  C3 and IGFBP5 log2FC in 11 cortical regions
#   igfbp5_across_layers.tsv        IGFBP5 across plasma, whole cortex, regions, other sets
#   igfbp5_celltypes.tsv     IGFBP5 by cell type, two snRNA sets, meta, LCM neurons
# Usage: python cortex_result_tables.py
# Where only a P value is available (Gandal 2018 microarray, plasma TMT), SE = |log2FC| / z with
# z = Phi^-1(1 - P/2) (normal approximation).
import csv
import math
import os

import openpyxl
from scipy.stats import norm

ROOT = os.environ.get("ASD_ROOT", ".")
BL = os.path.join(ROOT, 'brain')
OUT = os.path.join(BL, 'results')
PANEL_DIR = os.path.join(ROOT, 'panel12')
TAB, NL = chr(9), chr(10)
P12 = ['C3', 'AHSG', 'MBL2', 'C1RL', 'POSTN', 'BCHE', 'BTD', 'ANPEP', 'CLEC3B', 'IGFBP5',
       'PTGDS', 'QSOX1']
# row order of the disorder table (by site of synthesis)
P12_ORD = ['C3', 'AHSG', 'MBL2', 'C1RL', 'BCHE', 'BTD', 'ANPEP', 'IGFBP5', 'POSTN', 'CLEC3B',
           'QSOX1', 'PTGDS']
CT7 = ['EXN', 'INN', 'AST', 'ODC', 'OPC', 'MG', 'END']
REG = ['BA9', 'BA44_45', 'BA24', 'BA4_6', 'BA3_1_2_5', 'BA38', 'BA20_37', 'BA41_42_22',
       'BA39_40', 'BA7', 'BA17']
# Fatemi et al. 2024 Cereb Cortex, BA9, C3 log2FC reported in the main text (children, adults)
FATEMI_C3 = [('BA9, children', 0.9933), ('BA9, adults', 0.8051)]


def write(name, rows):
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, name), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, list(rows[0].keys()), delimiter=TAB, lineterminator=NL)
        w.writeheader()
        w.writerows(rows)
    print('wrote %-32s %4d rows' % (name, len(rows)))


def se_from_p(beta, p):
    z = norm.isf(p / 2) if p and p > 0 else float('nan')
    return abs(beta) / z if z and z > 0 else float('nan')


def rd(path):
    return list(csv.DictReader(open(path, encoding='utf-8'), delimiter=TAB))


def f(x):
    try:
        v = float(x)
        return v if math.isfinite(v) else None
    except (TypeError, ValueError):
        return None


# ---- detection of the 12 genes in brain cell types --------------------------
sc = rd(BL + '/data/lr_singlecell/panel12_by_celltype_both_datasets.tsv')
write('brain_cell_detection.tsv', [dict(dataset=r['dataset'], celltype=r['celltype'], gene=r['gene'],
                                     detect_frac=r['detect_frac'], mean_cpm=r['mean_cpm'],
                                     expressed=r['expressed']) for r in sc if r['gene'] in P12])
# % nuclei with >= 1 UMI, averaged over Wamsley and Velmeshev for the seven broad cell types;
# PTGDS from Velmeshev only (not captured in the Wamsley matrix); 0 where neither dataset has a value
rows = []
for g in P12:
    for ct in CT7:
        v = [f(r['detect_frac']) for r in sc if r['gene'] == g and r['celltype'] == ct
             and not (r['dataset'] == 'Wamsley' and g == 'PTGDS')]
        v = [x for x in v if x is not None]
        rows.append(dict(gene=g, celltype=ct, n_datasets=len(v),
                         detect_frac_mean=sum(v) / len(v) if v else 0.0))
write('brain_cell_detection_mean.tsv', rows)

# ---- Gandal 2022 (RNA-seq) and Gandal 2018 (microarray) gene tables; first row per symbol ----
wb = openpyxl.load_workbook(BL + '/data/asd_cortex/Gandal2022_MOESM5.xlsx', read_only=True)
it = wb['DEGene_Statistics'].iter_rows(values_only=True)
h = list(next(it))
ix = {k: i for i, k in enumerate(h)}
G22 = {}
for r in it:
    s = r[ix['external_gene_name']]
    if s and s not in G22:
        G22[s] = r
wb.close()
T18 = {}
for r in csv.DictReader(open(BL + '/data/replication/gandal2018a/TableS1_Microarray_MetaAnalyses.csv',
                             encoding='utf-8')):
    s = r['hgnc_symbol'] or r['external_gene_id']
    if s and s not in T18:
        T18[s] = r

# ---- complement cascade genes and layers ----------------------------------------
cg = rd(BL + '/data/asd_cortex/gandal2022_complement_genes.tsv')
rows = []
for r in cg:
    t = T18.get(r['gene'])
    rows.append(dict(gene=r['gene'], layer=r['layer'], rnaseq_logFC=r['wc_logFC'], rnaseq_FDR=r['wc_FDR'],
                     regions_up=r['regions_up'],
                     microarray_logFC=f(t['ASD.beta_log2FC']) if t else '',
                     microarray_FDR=f(t['ASD.FDR']) if t else ''))
write('complement_cascade_genes.tsv', rows)
mods = {r['module']: r for r in rd(BL + '/data/asd_cortex/gandal2022_modules.tsv')}
rep = {(r['dataset'], r['module']): r for r in rd(BL + '/data/replication/replication_modules.tsv')}
rows = []
for m, r in mods.items():
    if not (m.startswith('CMP_') or m == 'COMPLEMENT_all_layers'):
        continue
    q = rep.get(('Gandal2018a microarray ASD', m))
    rows.append(dict(layer=m, n=r['n'], rnaseq_median=r['median_logFC'], rnaseq_frac_up=r['frac_up'],
                     rnaseq_competitive_p=r['competitive_p'],
                     microarray_median=q['median_logFC'] if q else '',
                     microarray_competitive_p=q['competitive_p'] if q else ''))
write('complement_cascade_layers.tsv', rows)

# ---- the 12 genes and the complement module across disorders ----
rows = []
for g in P12_ORD:
    r = G22.get(g)
    b, q = (f(r[ix['WholeCortex_ASD_logFC']]), f(r[ix['WholeCortex_ASD_FDR']])) if r else (None, None)
    rows.append(dict(feature=g, dataset='Gandal 2022 RNA-seq', disorder='ASD', log2FC=b, se='', p='', fdr=q))
    t = T18.get(g)
    for dis in ['ASD', 'SCZ', 'BD', 'MDD', 'IBD']:
        if t is None or f(t[dis + '.beta_log2FC']) is None:
            rows.append(dict(feature=g, dataset='Gandal 2018a microarray', disorder=dis, log2FC='', se='',
                             p='', fdr=''))
            continue
        b, p = f(t[dis + '.beta_log2FC']), f(t[dis + '.P.value'])
        z = norm.isf(p / 2) if p and p > 0 else None
        rows.append(dict(feature=g, dataset='Gandal 2018a microarray', disorder=dis, log2FC=b,
                         se=abs(b) / z if z else '', p=p, fdr=f(t[dis + '.FDR'])))
m22 = mods['COMPLEMENT_all_layers']
rows.append(dict(feature='Complement (module)', dataset='Gandal 2022 RNA-seq', disorder='ASD',
                 log2FC=float(m22['median_logFC']), se='', p=float(m22['competitive_p']), fdr=''))
for dis in ['ASD', 'SCZ', 'BD', 'MDD', 'IBD']:
    key = 'Gandal2018a IBD (bowel, non-brain control)' if dis == 'IBD' else 'Gandal2018a microarray ' + dis
    q = rep[(key, 'COMPLEMENT_all_layers')]
    rows.append(dict(feature='Complement (module)', dataset='Gandal 2018a microarray', disorder=dis,
                     log2FC=float(q['median_logFC']), se='', p=float(q['competitive_p']), fdr=''))
write('disorders_panel12.tsv', rows)

# ---- plasma TMT, Zhang 2023 and single-nucleus look-ups -----------------------------------
PL = {r['Gene']: r for r in csv.DictReader(open(PANEL_DIR + '/dia_validation/stats_dep_concordance.csv',
                                                  encoding='utf-8'))}


def plasma(g):
    b, p = float(PL[g]['tmt_log2fc']), float(PL[g]['tmt_p'])
    return b, se_from_p(b, p), p


def zhang(fn, g):
    wbz = openpyxl.load_workbook(BL + '/data/replication/zhang2023/' + fn, read_only=True)
    itz = wbz.worksheets[0].iter_rows(values_only=True)
    hz = list(next(itz))
    iz = {k: i for i, k in enumerate(hz) if k}
    for rz in itz:
        if rz[iz['genenames']] == g:
            return rz[iz['log2FoldChange']], rz[iz['lfcSE']], rz[iz['pvalue']], rz[iz['padj']]
    return None


de = rd(BL + '/data/lr_singlecell/de_genes_both_datasets.tsv')


def sn(ds, ct, g):
    # limma-voom logFC, SE = |logFC / t|, P, FDR
    for x in de:
        if x['dataset'] == ds and x['celltype'] == ct and x['gene'] == g:
            b = float(x['logFC'])
            return b, abs(b / float(x['t'])), float(x['P.Value']), float(x['adj.P.Val'])
    return None


def add(rows_, layer, source, detail, b, se, p, fdr, n):
    rows_.append(dict(layer=layer, source=source, detail=detail, log2FC=b, se=se, p=p, fdr=fdr, n=n))


# ---- C3 across layers ------------------------------------------------------------
rows = []
b, se, p = plasma('C3')
add(rows, 'Plasma protein', 'This study (TMT)', 'ASD vs non-ASD', b, se, p, '', '')
r = G22['C3']
add(rows, 'Cortex RNA', 'Gandal 2022', '11 regions, 49 / 54', r[ix['WholeCortex_ASD_logFC']], '', '',
    r[ix['WholeCortex_ASD_FDR']], '49/54')
t = T18['C3']
b, p = f(t['ASD.beta_log2FC']), f(t['ASD.P.value'])
add(rows, 'Cortex RNA', 'Gandal 2018a', 'microarray (same brains)', b, se_from_p(b, p), p, f(t['ASD.FDR']), '')
z = zhang('pnas.2206758120.sd01.xlsx', 'C3')
add(rows, 'Cortex RNA', 'Zhang 2023', 'superior temporal gyrus', z[0], z[1], z[2], z[3], '27/32')
ab = rd(BL + '/data/brain_proteome/abraham2019_protein_table.tsv')
ab = [x for x in ab if x['gene'] == 'C3'][0]
# Abraham 2019 reports ASD abundance as % of control -> log2(val / 100)
add(rows, 'Cortex protein', 'Abraham 2019', 'BA19', math.log2(float(ab['val']) / 100), '',
    float(ab['p']), float(ab['fdr']), '')
for detail, lfc in FATEMI_C3:
    add(rows, 'Cortex protein', 'Fatemi 2024', detail, lfc, '', '', '<1e-4', '5/5')
for ds, n in [('Wamsley', '25/28'), ('Velmeshev', '11/12')]:
    s = sn(ds, 'MG', 'C3')
    add(rows, 'Microglia RNA', ds + (' 2024' if ds == 'Wamsley' else ' 2019'), 'single nucleus',
        s[0], s[1], s[2], s[3], n)
write('c3_across_layers.tsv', rows)

# ---- Gandal 2022 region-level log2FC for C3 and IGFBP5 (rows in sheet order) ---------------
rows = []
for g in G22:
    if g in ['C3', 'IGFBP5']:
        r = G22[g]
        for reg in REG:
            rows.append(dict(gene=g, region=reg, log2FC=r[ix['ASD_%s_logFC' % reg]],
                             fdr=r[ix['ASD_%s_FDR' % reg]]))
write('gandal2022_regions.tsv', rows)

# ---- IGFBP5 across layers and regions --------------------------------------
rows = []
b, se, p = plasma('IGFBP5')
add(rows, 'Plasma protein', 'This study (TMT)', 'ASD vs non-ASD', b, se, p, '', '')
r = G22['IGFBP5']
add(rows, 'Cortex RNA', 'Gandal 2022', 'whole cortex', r[ix['WholeCortex_ASD_logFC']], '', '',
    r[ix['WholeCortex_ASD_FDR']], '49/54')
for reg in REG:
    add(rows, 'Cortex RNA, by region', 'Gandal 2022', reg, r[ix['ASD_%s_logFC' % reg]], '', '',
        r[ix['ASD_%s_FDR' % reg]], '')
t = T18['IGFBP5']
b, p = f(t['ASD.beta_log2FC']), f(t['ASD.P.value'])
add(rows, 'Cortex RNA, other sets', 'Gandal 2018a', 'microarray (same brains)', b, se_from_p(b, p), p,
    f(t['ASD.FDR']), '')
z = zhang('pnas.2206758120.sd01.xlsx', 'IGFBP5')
add(rows, 'Cortex RNA, other sets', 'Zhang 2023', 'STG bulk', z[0], z[1], z[2], z[3], '27/32')
write('igfbp5_across_layers.tsv', rows)

# ---- IGFBP5 by cell type ---------------------------------------------------
meta = {(r['gene'], r['celltype']): r for r in rd(BL + '/data/lr_singlecell/panel12_meta_two_datasets.tsv')}
exp = {(r['dataset'], r['celltype']): r['expressed'] for r in sc if r['gene'] == 'IGFBP5'}
rows = []
for ct in CT7:
    for ds in ['Wamsley', 'Velmeshev']:
        s = sn(ds, ct, 'IGFBP5')
        rows.append(dict(celltype=ct, source=ds, log2FC=s[0] if s else '', se=s[1] if s else '',
                         p=s[2] if s else '', expressed=exp.get((ds, ct), 'FALSE')))
    m = meta.get(('IGFBP5', ct))
    if m and m['meta_P'] != '':
        # fixed-effect meta SE from the two dataset SEs
        sw = [x for x in rows if x['celltype'] == ct and x['source'] == 'Wamsley'][0]['se']
        sv = [x for x in rows if x['celltype'] == ct and x['source'] == 'Velmeshev'][0]['se']
        mse = math.sqrt(1 / (1 / sw ** 2 + 1 / sv ** 2))
        rows.append(dict(celltype=ct, source='Meta (Wamsley + Velmeshev)', log2FC=float(m['meta_logFC']),
                         se=mse, p=float(m['meta_P']), expressed='TRUE'))
z = zhang('pnas.2206758120.sd06.xlsx', 'IGFBP5')
rows.append(dict(celltype='LCM_NEURON', source='Zhang 2023 (LCM neurons, STG)', log2FC=z[0], se=z[1],
                 p=z[2], expressed='TRUE'))
write('igfbp5_celltypes.tsv', rows)
